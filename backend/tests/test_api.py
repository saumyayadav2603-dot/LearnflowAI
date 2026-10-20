"""Smoke and regression tests for the existing upload/extract API.

All files use pytest's isolated tmp_path; no cloud credentials are required.
"""
from io import BytesIO

import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

import backend.main as backend_main


@pytest.fixture()
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(backend_main, "UPLOAD_DIR", tmp_path)
    monkeypatch.setattr(
        backend_main,
        "upload_file_to_s3",
        lambda local_file_path, object_key, content_type=None: (
            f"s3://test-bucket/{object_key}"
        ),
    )
    monkeypatch.setattr(
        backend_main,
        "start_knowledge_base_sync",
        lambda: {"ingestionJobId": "mock-ingestion-job"},
    )
    monkeypatch.setattr(
        backend_main.app.state,
        "chat_service",
        object(),
        raising=False,
    )
    with TestClient(backend_main.app) as api:
        yield api


def make_pdf():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "LearnFlow lesson notes")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def make_pptx():
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = "Machine Learning Basics"
    stream = BytesIO()
    slides.save(stream)
    return stream.getvalue()


def upload(client, name, content):
    return client.post("/materials/upload", files={"file": (name, content)})


def test_health_check(client):
    assert client.get("/health").json() == {"status": "ok"}


def test_upload_and_extract_pdf_with_page_reference(client):
    response = upload(client, "lecture.pdf", make_pdf())
    assert response.status_code == 200
    payload = response.json()
    identifier = payload["stored_filename"]
    assert identifier.endswith(".pdf")
    assert payload["s3_uri"] == f"s3://test-bucket/{identifier}"
    assert payload["ingestion_job_id"] == "mock-ingestion-job"
    result = client.post(f"/materials/{identifier}/extract")
    assert result.status_code == 200
    assert result.json()["content"][0]["reference_type"] == "page"
    assert result.json()["content"][0]["reference_number"] == 1
    assert "LearnFlow lesson notes" in result.json()["content"][0]["text"]


def test_upload_and_extract_pptx_with_slide_reference(client):
    response = upload(client, "lecture.pptx", make_pptx())
    assert response.status_code == 200
    payload = response.json()
    identifier = payload["stored_filename"]
    assert identifier.endswith(".pptx")
    assert payload["s3_uri"] == f"s3://test-bucket/{identifier}"
    assert payload["ingestion_job_id"] == "mock-ingestion-job"
    result = client.post(f"/materials/{identifier}/extract")
    assert result.status_code == 200
    assert result.json()["content"][0]["reference_type"] == "slide"
    assert result.json()["content"][0]["reference_number"] == 1
    assert "Machine Learning Basics" in result.json()["content"][0]["text"]


def test_unsupported_upload_type_is_rejected(client, tmp_path):
    response = upload(client, "malware.exe", b"not a supported type")
    assert response.status_code == 415
    assert not list(tmp_path.iterdir())


def test_upload_size_limit_and_cleanup(client, tmp_path, monkeypatch):
    monkeypatch.setattr(backend_main, "MAX_UPLOAD_SIZE_BYTES", 8)
    response = upload(client, "too-large.pdf", b"more than eight bytes")
    assert response.status_code == 413
    assert not list(tmp_path.iterdir())


def test_missing_document_returns_404(client):
    identifier = "a" * 32 + ".pdf"
    assert client.post(f"/materials/{identifier}/extract").status_code == 404


@pytest.mark.parametrize("identifier", ["invalid.pdf", "not-a-uuid", "a" * 32 + ".exe"])
def test_invalid_document_identifier(client, identifier):
    assert client.post(f"/materials/{identifier}/extract").status_code == 400


def test_malformed_pdf_is_rejected_during_extraction(client):
    response = upload(client, "broken.pdf", b"not really a PDF")
    if response.status_code == 415:
        # Hardened content validation may reject a forged PDF at upload time.
        return
    assert response.status_code == 200
    identifier = response.json()["stored_filename"]
    result = client.post(f"/materials/{identifier}/extract")
    assert result.status_code == 422


def test_legacy_ppt_has_explicit_conversion_error(client):
    response = upload(client, "legacy.ppt", b"old PowerPoint content")
    if response.status_code == 415:
        # A stricter upload validator can reject the synthetic invalid PPT.
        return
    assert response.status_code == 200
    result = client.post(f"/materials/{response.json()['stored_filename']}/extract")
    assert result.status_code == 415
    assert "pptx" in result.json()["detail"]
