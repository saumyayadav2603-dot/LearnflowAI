"""Upload content-check regressions; no network or cloud dependencies."""
from io import BytesIO
import pytest
import pymupdf
from pptx import Presentation
from fastapi.testclient import TestClient
import backend.main as api


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "UPLOAD_DIR", tmp_path)
    monkeypatch.setattr(
        api,
        "upload_file_to_s3",
        lambda local_file_path, object_key, content_type=None: (
            f"s3://test-bucket/{object_key}"
        ),
    )
    monkeypatch.setattr(
        api,
        "start_knowledge_base_sync",
        lambda: {"ingestionJobId": "mock-ingestion-job"},
    )
    return TestClient(api.app), tmp_path


def post(client, filename, data):
    return client.post("/materials/upload", files={"file": (filename, data)})


def pdf_bytes():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Verified lesson PDF")
    raw = doc.tobytes()
    doc.close()
    return raw


def pptx_bytes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slide with citation"
    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("filename,data", [
    ("fake.pdf", b"not a pdf"),
    ("fake.pptx", b"not a zip"),
    ("fake.ppt", b"not OLE"),
    ("fake.avi", b"RIFF0000WAVE"),
    ("fake.mp4", b"broken ftyp"),
    ("fake.mov", b"invalid"),
    ("fake.mkv", b"invalid"),
    ("fake.webm", b"invalid"),
])
def test_rejects_fake_upload_and_deletes_it(client, filename, data):
    api_client, directory = client
    response = post(api_client, filename, data)
    assert response.status_code == 415, response.text
    assert list(directory.iterdir()) == []


def test_rejects_empty_upload(client):
    api_client, directory = client
    response = post(api_client, "empty.pdf", b"")
    assert response.status_code == 422
    assert list(directory.iterdir()) == []


@pytest.mark.parametrize("ext,contents,kind", [
    ("pdf", pdf_bytes, "page"),
    ("pptx", pptx_bytes, "slide"),
])
def test_upload_and_extract_real_content(client, ext, contents, kind):
    api_client, directory = client
    response = post(api_client, f"lesson.{ext}", contents())
    assert response.status_code == 200, response.text
    payload = response.json()
    source = payload["stored_filename"]
    assert (directory / source).is_file()
    assert payload["s3_uri"] == f"s3://test-bucket/{source}"
    assert payload["ingestion_job_id"] == "mock-ingestion-job"
    output = api_client.post(f"/materials/{source}/extract")
    assert output.status_code == 200, output.text
    assert output.json()["content"][0]["reference_type"] == kind
    assert output.json()["content"][0]["reference_number"] == 1
